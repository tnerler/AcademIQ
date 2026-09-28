"""
.docx CV -> duz metin (LLM cikarimi icin). Belge sirasi korunur:
  - paragraflar satir satir,
  - veri tablolari (her hucre tek satir) 'Yil | Derece | Universite' gibi tek satira,
  - duzen tablolari (yan panel, iki sutun: hucrelerde birden fazla paragraf ya da ic tablo)
    hucre hucre, her hucrenin icerigi ayri blok olarak.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph


def docx_text(path: Path) -> str:
    return "\n".join(line for line in _blocks(Document(str(path))) if line.strip())


def _blocks(container) -> list[str]:
    lines: list[str] = []
    for item in container.iter_inner_content():
        if isinstance(item, Paragraph):
            lines.append(" ".join(item.text.split()))
        elif isinstance(item, Table):
            lines.extend(_table(item))
    return lines


def _table(table: Table) -> list[str]:
    lines: list[str] = []
    for row in table.rows:
        cells = _unique(row.cells)  # birlesik hucreler her sutunda tekrar doner
        if all(_is_simple(c) for c in cells):
            values = [" ".join(c.text.split()) for c in cells]
            if any(values):
                lines.append(" | ".join(v for v in values if v))
        else:
            for c in cells:
                lines.extend(_blocks(c))
                lines.append("")
    return lines


def _is_simple(cell: _Cell) -> bool:
    return not cell.tables and sum(1 for p in cell.paragraphs if p.text.strip()) <= 1


def _unique(cells: list[_Cell]) -> list[_Cell]:
    seen, out = set(), []
    for c in cells:
        if id(c._tc) not in seen:
            seen.add(id(c._tc))
            out.append(c)
    return out
